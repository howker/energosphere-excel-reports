"""Independent, read-only verification of generated passports against a source.

Requires openpyxl in the developer's Python; not required by the application.
"""
import argparse
from datetime import datetime, date, time
from pathlib import Path
import re
import warnings
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('source', type=Path)
parser.add_argument('output', type=Path)
args = parser.parse_args()
source = openpyxl.load_workbook(args.source, data_only=True)
sheet = source['Прил.1.1 (Сч,ТТ,ТН)']
anchors = {}
for merged in sheet.merged_cells.ranges:
    for row in range(merged.min_row, merged.max_row + 1):
        for col in range(merged.min_col, merged.max_col + 1):
            anchors[row, col] = (merged.min_row, merged.min_col)

def value(row, col):
    r, c = anchors.get((row, col), (row, col))
    cell = sheet.cell(r, c)
    val = cell.value
    if val is None:
        return ''
    if isinstance(val, (datetime, date)):
        return val.strftime('%d.%m.%Y')
    if isinstance(val, time):
        return format(openpyxl.utils.datetime.to_excel(val), '.15g')
    if isinstance(val, (int, float)) and re.fullmatch(r'0{2,}', cell.number_format):
        return f'{int(val):0{len(cell.number_format)}d}'
    return str(val).strip()

files = list(args.output.glob('*.xlsx'))
outputs = {}
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
for path in files:
    with zipfile.ZipFile(path) as archive:
        # Fast XML inspection avoids third-party loss of native drawing groups.
        refs = ET.fromstring(archive.read('xl/workbook.xml')).find('s:sheets', ns)
        rels = {x.get('Id'): x.get('Target') for x in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        cells = {}
        for sh in refs:
            if sh.get('name') not in ('Лист1', 'исх.данные'):
                continue
            rid = sh.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            xml = ET.fromstring(archive.read('xl/' + rels[rid]))
            cells[sh.get('name')] = {c.get('r'): ''.join(t.text or '' for t in c.findall('.//s:t', ns)) for c in xml.findall('.//s:sheetData/s:row/s:c', ns)}
        key = (cells['исх.данные'].get('C3'), cells['исх.данные'].get('E3'))
        assert key not in outputs, f'Duplicate connection: {key}'
        outputs[key] = cells
        assert 'xl/vbaProject.bin' not in archive.namelist()

count = 0
for row in range(10, sheet.max_row + 1):
    if not isinstance(sheet.cell(row, 1).value, int) or not sheet.cell(row, 5).value:
        continue
    key = (value(row, 3), value(row, 5))
    cells = outputs.pop(key)
    for offset in range(3):
        for col in range(1, 33):
            ref = f'{openpyxl.utils.get_column_letter(col)}{offset + 3}'
            assert cells['исх.данные'].get(ref, '') == value(row + offset, col), (row, ref)
    printed = cells['Лист1']
    expected = {'M5': value(row, 8), 'P5': value(row, 9), 'M7': value(row, 10), 'M12': value(row, 12), 'P12': value(row, 11), 'A18': value(row, 5)}
    for offset in range(3):
        ct = 28 + 6 * offset
        vt = 4 + 6 * offset
        for ref, col in [(f'L{ct}', 15), (f'N{ct}', 17), (f'P{ct}', 18), (f'L{ct+1}', 16), (f'N{ct+3}', 20), (f'R{vt}', 23), (f'T{vt}', 25), (f'V{vt}', 26), (f'R{vt+1}', 24), (f'T{vt+3}', 28)]:
            expected[ref] = value(row + offset, col)
    for ref, val in expected.items():
        assert printed.get(ref, '') == val, (row, ref, val, printed.get(ref))
    for ref in ('P7', 'P9', 'P10', 'N29', 'P29', 'T5', 'V5', 'W5', 'Y5', 'W21', 'B44', 'Q41', 'Q43'):
        assert not printed.get(ref), (row, 'stale example', ref)
    count += 1
assert not outputs, 'Unexpected outputs'
assert len(files) == count, (len(files), count)
print(f'PASS: {count} reports; source cells, printed equipment, dates, identifiers, blank unsupported fields')
